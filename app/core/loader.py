"""文档加载：从磁盘读取 .md / .txt（可选 .pdf / .docx）为纯文本。

可选依赖缺失时不会崩溃，而是抛出带安装提示的 UnsupportedFormatError，
由上层接口转成明确的 HTTP 错误信息。
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, Iterator, List, Tuple

from app.core.chunker import build_doc_id
from app.core.text_utils import normalize_text, sha256_text
from app.logging_conf import get_logger
from app.models import SourceDocument

logger = get_logger(__name__)

TEXT_SUFFIXES = {".md", ".markdown", ".txt"}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | PDF_SUFFIXES | DOCX_SUFFIXES

_TITLE_RE = re.compile(r"^\s{0,3}#\s+(.+?)\s*$", re.M)


class UnsupportedFormatError(RuntimeError):
    """文件后缀不受支持，或缺少对应的可选解析依赖。"""


def iter_document_files(root: Path, suffixes=SUPPORTED_SUFFIXES) -> Iterator[Path]:
    """递归遍历知识库目录，跳过隐藏文件与临时文件。"""
    if not root.exists():
        return
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        if path.name.startswith(".") or path.name.startswith("~$"):
            continue
        if path.suffix.lower() not in suffixes:
            continue
        yield path


def _read_text_file(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    # 最后兜底：忽略非法字节，保证不因为一个坏字符整篇失败
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError as exc:  # pragma: no cover - 取决于本地是否装了可选依赖
        raise UnsupportedFormatError(
            "解析 PDF 需要可选依赖，请执行： pip install pypdf"
        ) from exc

    reader = PdfReader(str(path))
    pages: List[str] = []
    for page_index, page in enumerate(reader.pages):
        try:
            content = page.extract_text() or ""
        except Exception as exc:  # 单页失败不影响整篇
            logger.warning("PDF 第 %d 页解析失败 path=%s err=%s", page_index + 1, path.name, exc)
            continue
        if content.strip():
            pages.append(content)
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    try:
        import docx  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedFormatError(
            "解析 DOCX 需要可选依赖，请执行： pip install python-docx"
        ) from exc

    document = docx.Document(str(path))
    blocks = [para.text for para in document.paragraphs]
    # 表格内容也一并取出（售后知识库里参数表很常见）
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append("| " + " | ".join(cells) + " |")
    return "\n".join(blocks)


def load_document(path: Path, root: Path) -> SourceDocument:
    """读取单个文件，返回带元数据的 SourceDocument。

    :param path: 文件绝对路径
    :param root: 知识库根目录，用于生成相对路径（引用来源标识）
    """
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedFormatError(
            f"不支持的文件类型 {suffix or '(无后缀)'}，当前支持：{sorted(SUPPORTED_SUFFIXES)}"
        )

    if suffix in TEXT_SUFFIXES:
        raw = _read_text_file(path)
    elif suffix in PDF_SUFFIXES:
        raw = _read_pdf(path)
    else:
        raw = _read_docx(path)

    text = normalize_text(raw)
    if not text:
        raise UnsupportedFormatError("文件内容为空，或未能提取出任何文本")

    relative = path.relative_to(root).as_posix()
    title = _extract_title(text) or path.stem

    metadata: Dict[str, object] = {
        "filename": path.name,
        "suffix": suffix,
        "size_bytes": path.stat().st_size,
    }
    return SourceDocument(
        doc_id=build_doc_id(relative),
        source=relative,
        title=title,
        text=text,
        checksum=sha256_text(text),
        metadata=metadata,
    )


def _extract_title(text: str) -> str:
    match = _TITLE_RE.search(text)
    if match:
        return match.group(1).strip()[:80]
    for line in text.split("\n"):
        stripped = line.strip()
        if stripped:
            return stripped[:80]
    return ""


def scan_documents(root: Path) -> Tuple[List[Path], List[str]]:
    """扫描目录，返回 (可处理文件列表, 被跳过的文件说明列表)。"""
    files: List[Path] = []
    skipped: List[str] = []
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if path.name.startswith(".") or path.name.startswith("~$"):
            continue
        if path.suffix.lower() in SUPPORTED_SUFFIXES:
            files.append(path)
        else:
            skipped.append(f"{path.relative_to(root).as_posix()}（不支持的类型 {path.suffix}）")
    return sorted(files), skipped
